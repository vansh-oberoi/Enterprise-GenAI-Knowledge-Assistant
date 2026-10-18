from pathlib import Path
from typing import Any
import io
import pandas as pd
import fitz
from docx import Document

SUPPORTED_EXTENSIONS = {'.pdf', '.docx', '.txt', '.md', '.csv'}

def extract_file(filename: str, content: bytes) -> list[dict[str, Any]]:
    """Extract text units while preserving page, table, or CSV-row provenance."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{suffix}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
    units: list[dict[str, Any]] = []
    if suffix == '.pdf':
        with fitz.open(stream=content, filetype='pdf') as pdf:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text('text').strip()
                if text:
                    units.append({'text': text, 'metadata': {'page': page_number}})
    elif suffix == '.docx':
        doc = Document(io.BytesIO(content))
        text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        if text:
            units.append({'text': text, 'metadata': {}})
        for table_number, table in enumerate(doc.tables, start=1):
            rows = [' | '.join(cell.text.strip() for cell in row.cells) for row in table.rows]
            table_text = '\n'.join(rows).strip()
            if table_text:
                units.append({'text': table_text, 'metadata': {'table': table_number}})
    elif suffix in {'.txt', '.md'}:
        text = content.decode('utf-8', errors='replace').strip()
        if text:
            units.append({'text': text, 'metadata': {}})
    elif suffix == '.csv':
        frame = pd.read_csv(io.BytesIO(content)).fillna('')
        for row_index, row in frame.iterrows():
            text = ' | '.join(f'{column}: {row[column]}' for column in frame.columns)
            if text.strip():
                units.append({'text': text, 'metadata': {'row': int(row_index) + 2}})
    return units

def chunk_units(filename: str, units: list[dict[str, Any]], chunk_size: int, overlap: int):
    if chunk_size < 100:
        raise ValueError('chunk_size must be at least 100 characters')
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError('chunk overlap must be >= 0 and smaller than chunk size')
    chunks = []
    for unit in units:
        text = unit['text']
        base_metadata = {'source': filename, **unit.get('metadata', {})}
        start = 0
        while start < len(text):
            end = min(len(text), start + chunk_size)
            chunk_text = text[start:end].strip()
            if chunk_text:
                chunks.append({'text': chunk_text, 'metadata': {**base_metadata, 'chunk_start': start}})
            if end >= len(text):
                break
            start = end - overlap
    return chunks
